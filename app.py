import csv
import hmac
import io
import math
import os
import re
import secrets
import sqlite3
from contextlib import contextmanager
from datetime import datetime, timedelta
from pathlib import Path
from urllib.parse import quote, urlparse
from zoneinfo import ZoneInfo

import qrcode
from flask import (
    Flask,
    Response,
    abort,
    flash,
    jsonify,
    redirect,
    render_template,
    request,
    send_file,
    session,
    url_for,
)
from openpyxl import load_workbook
from werkzeug.security import check_password_hash, generate_password_hash


BASE_DIR = Path(__file__).resolve().parent
DB_PATH = Path(os.environ.get("SIGNIN_DATABASE", BASE_DIR / "data" / "signin.db"))
MAX_PASSPHRASE_LENGTH = 12
LOCAL_TIMEZONE = ZoneInfo(os.environ.get("APP_TIMEZONE", "Asia/Shanghai"))


def now_text():
    return datetime.now(LOCAL_TIMEZONE).isoformat(timespec="seconds")


def local_datetime(value):
    parsed = datetime.fromisoformat(value)
    if parsed.tzinfo is None:
        return parsed.replace(tzinfo=LOCAL_TIMEZONE)
    return parsed.astimezone(LOCAL_TIMEZONE)


def late_grace_minutes_from_form(form):
    value = str(form.get("late_grace_minutes", "0")).strip() or "0"
    try:
        minutes = int(value)
    except ValueError as error:
        raise ValueError("迟到宽限时间必须是整数分钟。") from error
    if not 0 <= minutes <= 1440:
        raise ValueError("迟到宽限时间必须在 0 到 1440 分钟之间。")
    return minutes


def attendance_status_for(activity, signed_at=None):
    if not activity["start_time"]:
        return "normal"
    try:
        start_time = local_datetime(activity["start_time"])
        check_time = local_datetime(signed_at) if isinstance(signed_at, str) else signed_at
        check_time = check_time or datetime.now(LOCAL_TIMEZONE)
        if check_time.tzinfo is None:
            check_time = check_time.replace(tzinfo=LOCAL_TIMEZONE)
        deadline = start_time + timedelta(minutes=activity["late_grace_minutes"] or 0)
        return "late" if check_time > deadline else "normal"
    except (TypeError, ValueError):
        return "normal"


def attendance_status_text(row):
    if not row["signed_in"]:
        return "未签到"
    labels = {
        "normal": "正常签到",
        "late": "迟到",
        "manual": "管理员补签",
    }
    return labels.get(row["attendance_status"] or "", "正常签到")


def new_credential_token(db):
    while True:
        token = secrets.token_urlsafe(24)
        exists = db.execute(
            "SELECT 1 FROM participants WHERE credential_token = ?", (token,)
        ).fetchone()
        if not exists:
            return token


def credential_token_from_value(value):
    text = str(value or "").strip()
    if not text or len(text) > 2048:
        return ""
    if re.fullmatch(r"[A-Za-z0-9_-]{24,128}", text):
        return text
    parsed = urlparse(text)
    path = parsed.path or text
    match = re.search(r"(?:^|/)credential/([A-Za-z0-9_-]{24,128})/?$", path)
    return match.group(1) if match else ""


def credential_payload(db, token, activity_id=None):
    params = [token]
    condition = "p.credential_token = ?"
    if activity_id is not None:
        condition += " AND p.activity_id = ?"
        params.append(activity_id)
    row = db.execute(
        f"""
        SELECT p.*, a.name AS activity_name, a.code AS activity_code,
               a.credential_check_enabled
        FROM participants p
        JOIN activities a ON a.id = p.activity_id
        WHERE {condition}
        """,
        params,
    ).fetchone()
    if row is None:
        return None
    return {
        "valid": bool(row["signed_in"] and row["credential_check_enabled"]),
        "verification_enabled": bool(row["credential_check_enabled"]),
        "name": row["name"],
        "identifier": row["identifier"],
        "shift": row["shift"],
        "position": row["position"],
        "sign_time": row["sign_time"],
        "attendance_status": row["attendance_status"],
        "attendance_status_text": attendance_status_text(row),
        "source": row["source"],
        "activity_name": row["activity_name"],
        "activity_code": row["activity_code"],
    }


def csv_safe(value):
    """Prevent spreadsheet formula execution when exported CSV is opened."""
    text = str(value or "")
    return "'" + text if text.startswith(("=", "+", "-", "@")) else text


def location_settings_from_form(form):
    enabled = str(form.get("location_check_enabled", "")).lower() in {
        "1",
        "true",
        "yes",
        "on",
    }
    latitude_text = str(form.get("latitude", "")).strip()
    longitude_text = str(form.get("longitude", "")).strip()
    radius_text = str(form.get("location_radius", "200")).strip() or "200"

    if enabled and (not latitude_text or not longitude_text):
        raise ValueError("开启位置核验时，请填写签到点的纬度和经度。")
    if bool(latitude_text) != bool(longitude_text):
        raise ValueError("纬度和经度需要同时填写。")

    latitude = longitude = None
    if latitude_text and longitude_text:
        try:
            latitude = float(latitude_text)
            longitude = float(longitude_text)
        except ValueError as error:
            raise ValueError("签到点经纬度格式不正确。") from error
        if not math.isfinite(latitude) or not -90 <= latitude <= 90:
            raise ValueError("纬度必须在 -90 到 90 之间。")
        if not math.isfinite(longitude) or not -180 <= longitude <= 180:
            raise ValueError("经度必须在 -180 到 180 之间。")

    try:
        radius = int(round(float(radius_text)))
    except ValueError as error:
        raise ValueError("位置核验半径必须是数字。") from error
    if not 10 <= radius <= 5000:
        raise ValueError("位置核验半径必须在 10 到 5000 米之间。")
    return int(enabled), latitude, longitude, radius


def passphrase_settings_from_form(form, current_activity=None):
    enabled = str(form.get("passphrase_check_enabled", "")).lower() in {
        "1",
        "true",
        "yes",
        "on",
    }
    if not enabled:
        return 0, "", 0

    passphrase = str(form.get("passphrase", ""))
    if not passphrase:
        if (
            current_activity is not None
            and current_activity["passphrase_check_enabled"]
            and current_activity["passphrase_hash"]
            and current_activity["passphrase_length"]
        ):
            return (
                1,
                current_activity["passphrase_hash"],
                current_activity["passphrase_length"],
            )
        raise ValueError("开启口令核验时，请设置签到口令。")
    if len(passphrase) > MAX_PASSPHRASE_LENGTH:
        raise ValueError(f"签到口令不能超过 {MAX_PASSPHRASE_LENGTH} 个字符。")
    if any(character.isspace() for character in passphrase):
        raise ValueError("签到口令不能包含空格或换行。")
    return 1, generate_password_hash(passphrase), len(passphrase)


def credential_check_enabled_from_form(form):
    return int(
        str(form.get("credential_check_enabled", "")).lower()
        in {"1", "true", "yes", "on"}
    )


def distance_in_meters(latitude_1, longitude_1, latitude_2, longitude_2):
    earth_radius = 6_371_000
    lat_1 = math.radians(latitude_1)
    lat_2 = math.radians(latitude_2)
    delta_lat = math.radians(latitude_2 - latitude_1)
    delta_lon = math.radians(longitude_2 - longitude_1)
    haversine = (
        math.sin(delta_lat / 2) ** 2
        + math.cos(lat_1) * math.cos(lat_2) * math.sin(delta_lon / 2) ** 2
    )
    return earth_radius * 2 * math.atan2(math.sqrt(haversine), math.sqrt(1 - haversine))


def submitted_location(form):
    try:
        latitude = float(form.get("latitude", ""))
        longitude = float(form.get("longitude", ""))
        accuracy = float(form.get("location_accuracy", ""))
    except (TypeError, ValueError) as error:
        raise ValueError("未取得有效位置，请重新授权定位后再签到。") from error
    if not math.isfinite(latitude) or not -90 <= latitude <= 90:
        raise ValueError("定位纬度无效，请重新定位。")
    if not math.isfinite(longitude) or not -180 <= longitude <= 180:
        raise ValueError("定位经度无效，请重新定位。")
    if not math.isfinite(accuracy) or not 0 <= accuracy <= 50_000:
        raise ValueError("定位精度无效，请重新定位。")
    return latitude, longitude, accuracy


def location_status_text(row):
    labels = {
        "verified": "范围内",
        "manual": "管理员补签",
        "not_required": "未启用",
    }
    return labels.get(row["location_status"] or "", "—")


@contextmanager
def get_db():
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    db = sqlite3.connect(DB_PATH, timeout=10)
    db.row_factory = sqlite3.Row
    db.execute("PRAGMA foreign_keys = ON")
    try:
        yield db
        db.commit()
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()


def init_db():
    with get_db() as db:
        db.executescript(
            """
            CREATE TABLE IF NOT EXISTS activities (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL,
                location TEXT NOT NULL DEFAULT '',
                description TEXT NOT NULL DEFAULT '',
                start_time TEXT,
                end_time TEXT,
                status TEXT NOT NULL DEFAULT 'draft'
                    CHECK(status IN ('draft', 'active', 'ended')),
                code TEXT NOT NULL UNIQUE,
                observer_token TEXT NOT NULL,
                location_check_enabled INTEGER NOT NULL DEFAULT 0,
                latitude REAL,
                longitude REAL,
                location_radius INTEGER NOT NULL DEFAULT 200,
                passphrase_check_enabled INTEGER NOT NULL DEFAULT 0,
                passphrase_hash TEXT NOT NULL DEFAULT '',
                passphrase_length INTEGER NOT NULL DEFAULT 0,
                late_grace_minutes INTEGER NOT NULL DEFAULT 0,
                credential_check_enabled INTEGER NOT NULL DEFAULT 0,
                created_at TEXT NOT NULL
            );

            CREATE TABLE IF NOT EXISTS participants (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                activity_id INTEGER NOT NULL REFERENCES activities(id) ON DELETE CASCADE,
                name TEXT NOT NULL,
                identifier TEXT NOT NULL,
                phone TEXT NOT NULL DEFAULT '',
                shift TEXT NOT NULL DEFAULT '',
                position TEXT NOT NULL DEFAULT '',
                signed_in INTEGER NOT NULL DEFAULT 0 CHECK(signed_in IN (0, 1)),
                sign_time TEXT,
                source TEXT NOT NULL DEFAULT '',
                note TEXT NOT NULL DEFAULT '',
                location_status TEXT NOT NULL DEFAULT '',
                location_distance REAL,
                location_accuracy REAL,
                attendance_status TEXT NOT NULL DEFAULT '',
                credential_token TEXT NOT NULL DEFAULT '',
                updated_at TEXT NOT NULL,
                UNIQUE(activity_id, identifier)
            );

            CREATE INDEX IF NOT EXISTS idx_participants_activity_signed
            ON participants(activity_id, signed_in);
            """
        )
        participant_columns = {
            row["name"] for row in db.execute("PRAGMA table_info(participants)").fetchall()
        }
        if "position" not in participant_columns:
            db.execute(
                "ALTER TABLE participants ADD COLUMN position TEXT NOT NULL DEFAULT ''"
            )
        participant_migrations = {
            "shift": "TEXT NOT NULL DEFAULT ''",
            "location_status": "TEXT NOT NULL DEFAULT ''",
            "location_distance": "REAL",
            "location_accuracy": "REAL",
            "attendance_status": "TEXT NOT NULL DEFAULT ''",
            "credential_token": "TEXT NOT NULL DEFAULT ''",
        }
        for column, definition in participant_migrations.items():
            if column not in participant_columns:
                db.execute(f"ALTER TABLE participants ADD COLUMN {column} {definition}")
        activity_columns = {
            row["name"] for row in db.execute("PRAGMA table_info(activities)").fetchall()
        }
        activity_migrations = {
            "observer_token": "TEXT NOT NULL DEFAULT ''",
            "location_check_enabled": "INTEGER NOT NULL DEFAULT 0",
            "latitude": "REAL",
            "longitude": "REAL",
            "location_radius": "INTEGER NOT NULL DEFAULT 200",
            "passphrase_check_enabled": "INTEGER NOT NULL DEFAULT 0",
            "passphrase_hash": "TEXT NOT NULL DEFAULT ''",
            "passphrase_length": "INTEGER NOT NULL DEFAULT 0",
            "late_grace_minutes": "INTEGER NOT NULL DEFAULT 0",
            "credential_check_enabled": "INTEGER NOT NULL DEFAULT 0",
        }
        for column, definition in activity_migrations.items():
            if column not in activity_columns:
                db.execute(f"ALTER TABLE activities ADD COLUMN {column} {definition}")
        used_tokens = set()
        activities = db.execute(
            "SELECT id, observer_token FROM activities ORDER BY id"
        ).fetchall()
        for activity in activities:
            token = activity["observer_token"] or ""
            if not token or token in used_tokens:
                token = secrets.token_urlsafe(24)
                while token in used_tokens:
                    token = secrets.token_urlsafe(24)
                db.execute(
                    "UPDATE activities SET observer_token = ? WHERE id = ?",
                    (token, activity["id"]),
                )
            used_tokens.add(token)
        db.execute(
            "CREATE UNIQUE INDEX IF NOT EXISTS idx_activities_observer_token "
            "ON activities(observer_token)"
        )
        legacy_signed_rows = db.execute(
            """
            SELECT p.id, p.signed_in, p.sign_time, p.source, p.attendance_status,
                   a.start_time, a.late_grace_minutes
            FROM participants p
            JOIN activities a ON a.id = p.activity_id
            WHERE p.signed_in = 1 AND p.attendance_status = ''
            """
        ).fetchall()
        for row in legacy_signed_rows:
            status = (
                "manual"
                if row["source"] == "管理员补签"
                else attendance_status_for(row, row["sign_time"])
            )
            db.execute(
                "UPDATE participants SET attendance_status = ? WHERE id = ?",
                (status, row["id"]),
            )
        signed_without_credentials = db.execute(
            """
            SELECT p.id FROM participants p
            JOIN activities a ON a.id = p.activity_id
            WHERE p.signed_in = 1 AND p.credential_token = ''
              AND a.credential_check_enabled = 1
            """
        ).fetchall()
        for row in signed_without_credentials:
            db.execute(
                "UPDATE participants SET credential_token = ? WHERE id = ?",
                (new_credential_token(db), row["id"]),
            )
        db.execute(
            "CREATE UNIQUE INDEX IF NOT EXISTS idx_participants_credential_token "
            "ON participants(credential_token) WHERE credential_token <> ''"
        )


app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "dev-change-me-" + secrets.token_hex(8))
app.config["MAX_CONTENT_LENGTH"] = 8 * 1024 * 1024
init_db()


PUBLIC_ENDPOINTS = {
    "sign_page",
    "submit_sign",
    "credential_page",
    "credential_qr",
    "observer_page",
    "observer_participants_api",
    "observer_export_results",
    "observer_verify_credential",
    "healthz",
    "static",
    "login",
}


@app.before_request
def require_admin_login():
    password = os.environ.get("ADMIN_PASSWORD", "")
    if not password or request.endpoint in PUBLIC_ENDPOINTS or session.get("admin_authenticated"):
        return None
    if request.path.startswith(("/api/", "/participants/")):
        return jsonify({"ok": False, "message": "登录已过期，请刷新页面重新登录。"}), 401
    return redirect(url_for("login", next=request.full_path))


@app.context_processor
def auth_context():
    return {"auth_enabled": bool(os.environ.get("ADMIN_PASSWORD"))}


def get_activity(activity_id):
    with get_db() as db:
        activity = db.execute("SELECT * FROM activities WHERE id = ?", (activity_id,)).fetchone()
    if activity is None:
        abort(404)
    return activity


def get_observer_activity(token):
    with get_db() as db:
        activity = db.execute(
            "SELECT * FROM activities WHERE observer_token = ?", (token,)
        ).fetchone()
    if activity is None:
        abort(404)
    return activity


def roster_rows(upload):
    suffix = Path(upload.filename or "").suffix.lower()
    data = upload.read()
    if suffix == ".csv":
        decoded = None
        for encoding in ("utf-8-sig", "gb18030"):
            try:
                decoded = data.decode(encoding)
                break
            except UnicodeDecodeError:
                pass
        if decoded is None:
            raise ValueError("CSV 编码无法识别，请使用 UTF-8 或 GB18030。")
        return list(csv.reader(io.StringIO(decoded)))
    if suffix == ".xlsx":
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sheet = workbook.active
        return [["" if cell is None else str(cell).strip() for cell in row] for row in sheet.iter_rows(values_only=True)]
    raise ValueError("仅支持 CSV 或 XLSX 文件。")


HEADER_ALIASES = {
    "name": {"姓名", "名字", "name"},
    "identifier": {"学号", "工号", "学号/工号", "编号", "id", "identifier"},
    "phone": {"手机号", "手机", "联系电话", "电话", "phone", "mobile"},
    "shift": {"班次", "班别", "时段", "shift", "session"},
    "position": {"岗位", "职位", "职务", "position", "role"},
}


def parse_roster(upload):
    rows = roster_rows(upload)
    if not rows:
        raise ValueError("名单为空。")
    headers = [str(value).strip().lower().replace(" ", "") for value in rows[0]]
    indexes = {}
    for key, aliases in HEADER_ALIASES.items():
        aliases_lower = {item.lower().replace(" ", "") for item in aliases}
        for index, header in enumerate(headers):
            if header in aliases_lower:
                indexes[key] = index
                break
    if "name" not in indexes or "identifier" not in indexes:
        raise ValueError("缺少必填表头：姓名、学号/工号。")

    parsed = []
    for line_number, row in enumerate(rows[1:], start=2):
        values = [str(value).strip() for value in row]
        name = values[indexes["name"]] if indexes["name"] < len(values) else ""
        identifier = values[indexes["identifier"]] if indexes["identifier"] < len(values) else ""
        phone_index = indexes.get("phone")
        phone = values[phone_index] if phone_index is not None and phone_index < len(values) else ""
        shift_index = indexes.get("shift")
        shift = values[shift_index] if shift_index is not None and shift_index < len(values) else ""
        position_index = indexes.get("position")
        position = (
            values[position_index]
            if position_index is not None and position_index < len(values)
            else ""
        )
        if not name and not identifier:
            continue
        if not name or not identifier:
            raise ValueError(f"第 {line_number} 行缺少姓名或学号/工号。")
        parsed.append((name, identifier, phone, shift, position))
    if not parsed:
        raise ValueError("名单中没有有效数据。")
    return parsed


@app.template_filter("display_time")
def display_time(value):
    if not value:
        return "—"
    try:
        return local_datetime(value).strftime("%m-%d %H:%M:%S")
    except ValueError:
        return value


@app.get("/")
def index():
    return redirect(url_for("admin"))


@app.route("/login", methods=["GET", "POST"])
def login():
    password = os.environ.get("ADMIN_PASSWORD", "")
    if not password:
        return redirect(url_for("admin"))
    if request.method == "POST":
        supplied = request.form.get("password", "")
        if hmac.compare_digest(supplied, password):
            session.clear()
            session["admin_authenticated"] = True
            target = request.form.get("next", "")
            if not target.startswith("/") or target.startswith("//"):
                target = url_for("admin")
            return redirect(target)
        flash("管理员密码不正确。", "danger")
    return render_template("login.html", next=request.args.get("next", ""), public_page=True)


@app.post("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


@app.get("/healthz")
def healthz():
    with get_db() as db:
        db.execute("SELECT 1").fetchone()
    return jsonify({"ok": True})


@app.get("/admin")
def admin():
    selected_id = request.args.get("activity_id", type=int)
    with get_db() as db:
        activities = db.execute(
            """
            SELECT a.*, COUNT(p.id) AS total,
                   SUM(CASE WHEN p.signed_in = 1 THEN 1 ELSE 0 END) AS signed
            FROM activities a LEFT JOIN participants p ON p.activity_id = a.id
            GROUP BY a.id ORDER BY a.created_at DESC
            """
        ).fetchall()
        if selected_id:
            activity = db.execute("SELECT * FROM activities WHERE id = ?", (selected_id,)).fetchone()
        else:
            activity = db.execute(
                "SELECT * FROM activities ORDER BY created_at DESC LIMIT 1"
            ).fetchone()
    return render_template("admin.html", activities=activities, activity=activity)


@app.post("/activities")
def create_activity():
    name = request.form.get("name", "").strip()
    if not name:
        flash("请填写活动名称。", "danger")
        return redirect(url_for("admin"))
    try:
        location_check_enabled, latitude, longitude, location_radius = (
            location_settings_from_form(request.form)
        )
        passphrase_check_enabled, passphrase_hash, passphrase_length = (
            passphrase_settings_from_form(request.form)
        )
        late_grace_minutes = late_grace_minutes_from_form(request.form)
        credential_check_enabled = credential_check_enabled_from_form(request.form)
    except ValueError as error:
        flash(str(error), "danger")
        return redirect(url_for("admin"))
    code = secrets.token_urlsafe(6).replace("-", "").replace("_", "")[:8].upper()
    observer_token = secrets.token_urlsafe(24)
    with get_db() as db:
        cursor = db.execute(
            """
            INSERT INTO activities(
                name, location, description, start_time, end_time, status,
                code, observer_token, location_check_enabled, latitude, longitude,
                location_radius, passphrase_check_enabled, passphrase_hash,
                passphrase_length, late_grace_minutes, credential_check_enabled,
                created_at
            )
            VALUES (?, ?, ?, ?, ?, 'draft', ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                name,
                request.form.get("location", "").strip(),
                request.form.get("description", "").strip(),
                request.form.get("start_time") or None,
                request.form.get("end_time") or None,
                code,
                observer_token,
                location_check_enabled,
                latitude,
                longitude,
                location_radius,
                passphrase_check_enabled,
                passphrase_hash,
                passphrase_length,
                late_grace_minutes,
                credential_check_enabled,
                now_text(),
            ),
        )
        activity_id = cursor.lastrowid
    flash("活动已创建，请上传名单并启动活动。", "success")
    return redirect(url_for("admin", activity_id=activity_id))


@app.post("/activities/<int:activity_id>/status")
def set_activity_status(activity_id):
    get_activity(activity_id)
    status = request.form.get("status")
    if status not in {"draft", "active", "ended"}:
        abort(400)
    with get_db() as db:
        db.execute("UPDATE activities SET status = ? WHERE id = ?", (status, activity_id))
    labels = {"draft": "准备中", "active": "进行中", "ended": "已结束"}
    flash(f"活动状态已更新为“{labels[status]}”。", "success")
    return redirect(url_for("admin", activity_id=activity_id))


@app.post("/activities/<int:activity_id>/edit")
def edit_activity(activity_id):
    activity = get_activity(activity_id)
    name = request.form.get("name", "").strip()
    start_time = request.form.get("start_time") or None
    end_time = request.form.get("end_time") or None
    if not name:
        flash("活动名称不能为空。", "danger")
        return redirect(url_for("admin", activity_id=activity_id))
    try:
        if start_time:
            datetime.fromisoformat(start_time)
        if end_time:
            datetime.fromisoformat(end_time)
        if start_time and end_time and datetime.fromisoformat(end_time) < datetime.fromisoformat(start_time):
            raise ValueError("结束时间不能早于开始时间。")
    except ValueError as error:
        message = str(error) if str(error) == "结束时间不能早于开始时间。" else "活动时间格式不正确。"
        flash(message, "danger")
        return redirect(url_for("admin", activity_id=activity_id))
    try:
        location_check_enabled, latitude, longitude, location_radius = (
            location_settings_from_form(request.form)
        )
        passphrase_check_enabled, passphrase_hash, passphrase_length = (
            passphrase_settings_from_form(request.form, activity)
        )
        late_grace_minutes = late_grace_minutes_from_form(request.form)
        credential_check_enabled = credential_check_enabled_from_form(request.form)
    except ValueError as error:
        flash(str(error), "danger")
        return redirect(url_for("admin", activity_id=activity_id))
    with get_db() as db:
        db.execute(
            """
            UPDATE activities
            SET name = ?, location = ?, description = ?, start_time = ?, end_time = ?,
                location_check_enabled = ?, latitude = ?, longitude = ?, location_radius = ?,
                passphrase_check_enabled = ?, passphrase_hash = ?, passphrase_length = ?,
                late_grace_minutes = ?, credential_check_enabled = ?
            WHERE id = ?
            """,
            (
                name,
                request.form.get("location", "").strip(),
                request.form.get("description", "").strip(),
                start_time,
                end_time,
                location_check_enabled,
                latitude,
                longitude,
                location_radius,
                passphrase_check_enabled,
                passphrase_hash,
                passphrase_length,
                late_grace_minutes,
                credential_check_enabled,
                activity_id,
            ),
        )
        if credential_check_enabled:
            missing_credentials = db.execute(
                """
                SELECT id FROM participants
                WHERE activity_id = ? AND signed_in = 1 AND credential_token = ''
                """,
                (activity_id,),
            ).fetchall()
            for participant in missing_credentials:
                db.execute(
                    "UPDATE participants SET credential_token = ? WHERE id = ?",
                    (new_credential_token(db), participant["id"]),
                )
    flash("活动信息已更新。", "success")
    return redirect(url_for("admin", activity_id=activity_id))


@app.post("/activities/<int:activity_id>/upload")
def upload_roster(activity_id):
    get_activity(activity_id)
    upload = request.files.get("roster")
    if not upload or not upload.filename:
        flash("请选择 CSV 或 XLSX 名单。", "danger")
        return redirect(url_for("admin", activity_id=activity_id))
    try:
        rows = parse_roster(upload)
        with get_db() as db:
            for name, identifier, phone, shift, position in rows:
                db.execute(
                    """
                    INSERT INTO participants(
                        activity_id, name, identifier, phone, shift, position, updated_at
                    )
                    VALUES (?, ?, ?, ?, ?, ?, ?)
                    ON CONFLICT(activity_id, identifier) DO UPDATE SET
                        name = excluded.name,
                        phone = excluded.phone,
                        shift = excluded.shift,
                        position = excluded.position,
                        updated_at = excluded.updated_at
                    """,
                    (activity_id, name, identifier, phone, shift, position, now_text()),
                )
        flash(f"名单导入完成，共处理 {len(rows)} 人；重复编号已更新，签到状态予以保留。", "success")
    except (ValueError, OSError) as error:
        flash(str(error), "danger")
    return redirect(url_for("admin", activity_id=activity_id))


@app.get("/sign/<code>")
def sign_page(code):
    with get_db() as db:
        activity = db.execute("SELECT * FROM activities WHERE code = ?", (code.upper(),)).fetchone()
    if activity is None:
        abort(404)
    return render_template("sign.html", activity=activity, public_page=True)


@app.post("/sign/<code>")
def submit_sign(code):
    identifier = request.form.get("identifier", "").strip()
    name = request.form.get("name", "").strip()
    phone = request.form.get("phone", "").strip()
    note = request.form.get("note", "").strip()[:200]
    credential_token_to_show = None
    with get_db() as db:
        activity = db.execute("SELECT * FROM activities WHERE code = ?", (code.upper(),)).fetchone()
        if activity is None:
            abort(404)
        if activity["status"] != "active":
            flash("当前活动尚未开始或已经结束，暂不能签到。", "danger")
            return redirect(url_for("sign_page", code=code))
        participant = db.execute(
            "SELECT * FROM participants WHERE activity_id = ? AND identifier = ?",
            (activity["id"], identifier),
        ).fetchone()
        if participant is None:
            flash("名单中未找到该学号/工号，请联系现场管理员。", "danger")
        elif participant["name"].strip().lower() != name.lower():
            flash("姓名与名单不匹配，请检查后重试。", "danger")
        elif participant["phone"] and phone and participant["phone"] != phone:
            flash("手机号与名单不匹配，请检查后重试。", "danger")
        elif participant["signed_in"]:
            flash(f"{participant['name']} 已完成签到，无需重复操作。", "info")
            if activity["credential_check_enabled"]:
                credential_token_to_show = participant["credential_token"] or new_credential_token(db)
                if not participant["credential_token"]:
                    db.execute(
                        "UPDATE participants SET credential_token = ?, updated_at = ? WHERE id = ?",
                        (credential_token_to_show, now_text(), participant["id"]),
                    )
        else:
            location_status = "not_required"
            location_distance = None
            location_accuracy = None
            if activity["location_check_enabled"]:
                try:
                    latitude, longitude, location_accuracy = submitted_location(request.form)
                except ValueError as error:
                    flash(str(error), "danger")
                    return redirect(url_for("sign_page", code=code))
                location_distance = distance_in_meters(
                    activity["latitude"], activity["longitude"], latitude, longitude
                )
                if location_distance > activity["location_radius"]:
                    flash(
                        f"当前位置距签到点约 {round(location_distance)} 米，"
                        f"超出允许的 {activity['location_radius']} 米范围。",
                        "danger",
                    )
                    return redirect(url_for("sign_page", code=code))
                location_status = "verified"
            if activity["passphrase_check_enabled"]:
                submitted_passphrase = str(request.form.get("passphrase", ""))
                if not submitted_passphrase or not check_password_hash(
                    activity["passphrase_hash"], submitted_passphrase
                ):
                    flash("签到口令不正确，请核对后重试。", "danger")
                    return redirect(url_for("sign_page", code=code))
            sign_time = now_text()
            attendance_status = attendance_status_for(activity, sign_time)
            credential_token = participant["credential_token"] or ""
            if activity["credential_check_enabled"] and not credential_token:
                credential_token = new_credential_token(db)
            cursor = db.execute(
                """
                UPDATE participants SET signed_in = 1, sign_time = ?, source = '扫码签到',
                    note = ?, location_status = ?, location_distance = ?,
                    location_accuracy = ?, attendance_status = ?, credential_token = ?,
                    updated_at = ?
                WHERE id = ? AND signed_in = 0
                """,
                (
                    sign_time,
                    note,
                    location_status,
                    location_distance,
                    location_accuracy,
                    attendance_status,
                    credential_token,
                    sign_time,
                    participant["id"],
                ),
            )
            if cursor.rowcount:
                suffix = "（已标记迟到）" if attendance_status == "late" else ""
                flash(f"签到成功，欢迎你，{participant['name']}！{suffix}", "success")
                if activity["credential_check_enabled"]:
                    credential_token_to_show = credential_token
    if credential_token_to_show:
        return redirect(url_for("credential_page", token=credential_token_to_show))
    return redirect(url_for("sign_page", code=code))


def qr_image_response(value, filename):
    qr = qrcode.QRCode(version=None, box_size=8, border=3)
    qr.add_data(value)
    qr.make(fit=True)
    image = qr.make_image(fill_color="#000000", back_color="white")
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    buffer.seek(0)
    return send_file(buffer, mimetype="image/png", download_name=filename)


@app.get("/credential/<token>")
def credential_page(token):
    with get_db() as db:
        credential = credential_payload(db, token)
    if credential is None or not credential["verification_enabled"]:
        abort(404)
    return render_template(
        "credential.html",
        credential=credential,
        credential_token=token,
        public_page=True,
    )


@app.get("/credential/<token>/qr.png")
def credential_qr(token):
    with get_db() as db:
        credential = credential_payload(db, token)
    if credential is None or not credential["verification_enabled"]:
        abort(404)
    public_base = os.environ.get("PUBLIC_BASE_URL", request.url_root.rstrip("/"))
    credential_url = f"{public_base}{url_for('credential_page', token=token)}"
    return qr_image_response(credential_url, f"credential-{credential['identifier']}.png")


def verify_credential_response(activity):
    if not activity["credential_check_enabled"]:
        return jsonify({"ok": False, "message": "当前活动未启用电子凭证二次核验。"}), 403
    submitted = request.get_json(silent=True) or request.form
    token = credential_token_from_value(submitted.get("value", ""))
    if not token:
        return jsonify({"ok": False, "message": "未识别到有效的签到凭证二维码。"}), 400
    with get_db() as db:
        credential = credential_payload(db, token, activity["id"])
    if credential is None:
        return jsonify({"ok": False, "message": "该凭证不属于当前活动或凭证不存在。"}), 404
    if not credential["valid"]:
        return jsonify(
            {"ok": False, "message": "该凭证已失效，人员当前处于未签到状态。", "credential": credential}
        ), 409
    return jsonify({"ok": True, "message": "凭证有效，二次核验通过。", "credential": credential})


@app.post("/api/activities/<int:activity_id>/credentials/verify")
def admin_verify_credential(activity_id):
    activity = get_activity(activity_id)
    return verify_credential_response(activity)


@app.post("/observe/<token>/credentials/verify")
def observer_verify_credential(token):
    activity = get_observer_activity(token)
    return verify_credential_response(activity)


@app.get("/activities/<int:activity_id>/qr.png")
def activity_qr(activity_id):
    activity = get_activity(activity_id)
    public_base = os.environ.get("PUBLIC_BASE_URL", request.url_root.rstrip("/"))
    sign_url = f"{public_base}/sign/{activity['code']}"
    return qr_image_response(sign_url, f"{activity['code']}.png")


@app.get("/api/activities/<int:activity_id>/participants")
def participants_api(activity_id):
    get_activity(activity_id)
    status = request.args.get("status", "all")
    search = request.args.get("q", "").strip()
    shift_filter = request.args.get("shift", "").strip()
    position_filter = request.args.get("position", "").strip()
    return jsonify(
        participant_payload(activity_id, status, search, shift_filter, position_filter)
    )


def group_summary(db, activity_id, field, empty_label):
    if field not in {"shift", "position"}:
        raise ValueError("Unsupported grouping field")
    rows = db.execute(
        f"""
        SELECT {field} AS group_value,
               COUNT(*) AS total,
               SUM(CASE WHEN signed_in = 1 THEN 1 ELSE 0 END) AS signed,
               SUM(CASE WHEN attendance_status = 'late' THEN 1 ELSE 0 END) AS late
        FROM participants
        WHERE activity_id = ?
        GROUP BY {field}
        ORDER BY CASE WHEN TRIM({field}) = '' THEN 1 ELSE 0 END,
                 {field} COLLATE NOCASE
        """,
        (activity_id,),
    ).fetchall()
    return [
        {
            "value": row["group_value"] or "",
            "label": row["group_value"] or empty_label,
            "total": row["total"],
            "signed": row["signed"] or 0,
            "unsigned": row["total"] - (row["signed"] or 0),
            "late": row["late"] or 0,
            "rate": round((row["signed"] or 0) * 100 / row["total"], 1),
        }
        for row in rows
    ]


def participant_payload(
    activity_id, status="all", search="", shift_filter="", position_filter=""
):
    conditions = ["activity_id = ?"]
    params = [activity_id]
    if status == "signed":
        conditions.append("signed_in = 1")
    elif status == "unsigned":
        conditions.append("signed_in = 0")
    elif status == "late":
        conditions.append("signed_in = 1 AND attendance_status = 'late'")
    elif status == "exception":
        conditions.append("signed_in = 1 AND attendance_status IN ('late', 'manual')")
    for column, value in (("shift", shift_filter), ("position", position_filter)):
        if value == "__empty__":
            conditions.append(f"TRIM({column}) = ''")
        elif value:
            conditions.append(f"{column} = ?")
            params.append(value)
    if search:
        conditions.append(
            "(name LIKE ? OR identifier LIKE ? OR phone LIKE ? OR shift LIKE ? OR position LIKE ?)"
        )
        term = f"%{search}%"
        params.extend([term, term, term, term, term])
    where = " AND ".join(conditions)
    with get_db() as db:
        counts = db.execute(
            """
            SELECT COUNT(*) AS total,
                   SUM(CASE WHEN signed_in = 1 THEN 1 ELSE 0 END) AS signed,
                   SUM(CASE WHEN attendance_status = 'late' THEN 1 ELSE 0 END) AS late,
                   SUM(CASE WHEN attendance_status IN ('late', 'manual') THEN 1 ELSE 0 END) AS exception_count
            FROM participants WHERE activity_id = ?
            """,
            (activity_id,),
        ).fetchone()
        rows = db.execute(
            f"SELECT * FROM participants WHERE {where} ORDER BY signed_in ASC, name COLLATE NOCASE ASC",
            params,
        ).fetchall()
        groups = {
            "shifts": group_summary(db, activity_id, "shift", "未分班次"),
            "positions": group_summary(db, activity_id, "position", "未分岗位"),
        }
    total = counts["total"] or 0
    signed = counts["signed"] or 0
    return {
        "summary": {
            "total": total,
            "signed": signed,
            "unsigned": total - signed,
            "rate": round(signed * 100 / total, 1) if total else 0,
            "late": counts["late"] or 0,
            "exception": counts["exception_count"] or 0,
        },
        "participants": [dict(row) for row in rows],
        "groups": groups,
    }


@app.get("/observe/<token>")
def observer_page(token):
    activity = get_observer_activity(token)
    return render_template("observer.html", activity=activity, public_page=True)


@app.get("/observe/<token>/participants")
def observer_participants_api(token):
    activity = get_observer_activity(token)
    status = request.args.get("status", "all")
    search = request.args.get("q", "").strip()
    shift_filter = request.args.get("shift", "").strip()
    position_filter = request.args.get("position", "").strip()
    payload = participant_payload(
        activity["id"], status, search, shift_filter, position_filter
    )
    visible_fields = (
        "name",
        "identifier",
        "phone",
        "shift",
        "position",
        "signed_in",
        "attendance_status",
        "sign_time",
        "source",
        "note",
        "location_status",
        "location_distance",
        "location_accuracy",
    )
    payload["participants"] = [
        {field: participant[field] for field in visible_fields}
        for participant in payload["participants"]
    ]
    return jsonify(payload)


@app.post("/participants/<int:participant_id>/toggle")
def toggle_participant(participant_id):
    payload = request.get_json(silent=True) or request.form
    signed = str(payload.get("signed", "")).lower() in {"1", "true", "yes"}
    note = str(payload.get("note", "")).strip()[:200]
    with get_db() as db:
        participant = db.execute("SELECT * FROM participants WHERE id = ?", (participant_id,)).fetchone()
        if participant is None:
            abort(404)
        if signed:
            activity = db.execute(
                "SELECT credential_check_enabled FROM activities WHERE id = ?",
                (participant["activity_id"],),
            ).fetchone()
            credential_token = participant["credential_token"] or ""
            if activity["credential_check_enabled"] and not credential_token:
                credential_token = new_credential_token(db)
            db.execute(
                """
                UPDATE participants SET signed_in = 1, sign_time = ?, source = '管理员补签',
                       note = ?, location_status = 'manual', location_distance = NULL,
                       location_accuracy = NULL, attendance_status = 'manual',
                       credential_token = ?, updated_at = ? WHERE id = ?
                """,
                (now_text(), note, credential_token, now_text(), participant_id),
            )
        else:
            db.execute(
                """
                UPDATE participants SET signed_in = 0, sign_time = NULL, source = '管理员取消',
                       note = ?, location_status = '', location_distance = NULL,
                       location_accuracy = NULL, attendance_status = '',
                       updated_at = ? WHERE id = ?
                """,
                (note, now_text(), participant_id),
            )
    return jsonify({"ok": True})


@app.patch("/participants/<int:participant_id>")
def edit_participant(participant_id):
    payload = request.get_json(silent=True) or {}
    name = str(payload.get("name", "")).strip()
    identifier = str(payload.get("identifier", "")).strip()
    phone = str(payload.get("phone", "")).strip()
    shift = str(payload.get("shift", "")).strip()[:100]
    position = str(payload.get("position", "")).strip()[:100]
    note = str(payload.get("note", "")).strip()[:200]
    if not name or not identifier:
        return jsonify({"ok": False, "message": "姓名和学号/工号不能为空。"}), 400
    with get_db() as db:
        participant = db.execute("SELECT * FROM participants WHERE id = ?", (participant_id,)).fetchone()
        if participant is None:
            abort(404)
        try:
            db.execute(
                """
                UPDATE participants
                SET name = ?, identifier = ?, phone = ?, shift = ?, position = ?,
                    note = ?, updated_at = ?
                WHERE id = ?
                """,
                (name, identifier, phone, shift, position, note, now_text(), participant_id),
            )
        except sqlite3.IntegrityError:
            return jsonify({"ok": False, "message": "该学号/工号已存在于本活动名单中。"}), 409
    return jsonify({"ok": True})


@app.delete("/participants/<int:participant_id>")
def delete_participant(participant_id):
    with get_db() as db:
        participant = db.execute("SELECT id FROM participants WHERE id = ?", (participant_id,)).fetchone()
        if participant is None:
            abort(404)
        db.execute("DELETE FROM participants WHERE id = ?", (participant_id,))
    return jsonify({"ok": True})


@app.get("/activities/<int:activity_id>/export.csv")
def export_results(activity_id):
    activity = get_activity(activity_id)
    status = request.args.get("status", "all")
    return export_response(
        activity,
        status,
        request.args.get("shift", "").strip(),
        request.args.get("position", "").strip(),
        request.args.get("q", "").strip(),
    )


@app.get("/observe/<token>/export.csv")
def observer_export_results(token):
    activity = get_observer_activity(token)
    status = request.args.get("status", "all")
    return export_response(
        activity,
        status,
        request.args.get("shift", "").strip(),
        request.args.get("position", "").strip(),
        request.args.get("q", "").strip(),
    )


def export_response(
    activity, status="all", shift_filter="", position_filter="", search=""
):
    conditions = ["activity_id = ?"]
    params = [activity["id"]]
    status_conditions = {
        "signed": "signed_in = 1",
        "unsigned": "signed_in = 0",
        "late": "signed_in = 1 AND attendance_status = 'late'",
        "exception": "signed_in = 1 AND attendance_status IN ('late', 'manual')",
    }
    if status in status_conditions:
        conditions.append(status_conditions[status])
    for column, value in (("shift", shift_filter), ("position", position_filter)):
        if value == "__empty__":
            conditions.append(f"TRIM({column}) = ''")
        elif value:
            conditions.append(f"{column} = ?")
            params.append(value)
    if search:
        conditions.append(
            "(name LIKE ? OR identifier LIKE ? OR phone LIKE ? OR shift LIKE ? OR position LIKE ?)"
        )
        term = f"%{search}%"
        params.extend([term, term, term, term, term])
    where = " AND ".join(conditions)
    with get_db() as db:
        rows = db.execute(
            f"SELECT * FROM participants WHERE {where} ORDER BY name COLLATE NOCASE",
            params,
        ).fetchall()
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(
        [
            "姓名",
            "学号/工号",
            "手机号",
            "班次",
            "岗位",
            "签到状态",
            "考勤状态",
            "签到时间",
            "位置核验",
            "距签到点(米)",
            "定位精度(米)",
            "来源",
            "备注",
        ]
    )
    for row in rows:
        writer.writerow(
            [
                csv_safe(row["name"]),
                csv_safe(row["identifier"]),
                csv_safe(row["phone"]),
                csv_safe(row["shift"]),
                csv_safe(row["position"]),
                "已签到" if row["signed_in"] else "未签到",
                attendance_status_text(row),
                row["sign_time"] or "",
                location_status_text(row),
                round(row["location_distance"]) if row["location_distance"] is not None else "",
                round(row["location_accuracy"]) if row["location_accuracy"] is not None else "",
                csv_safe(row["source"]),
                csv_safe(row["note"]),
            ]
        )
    filename = f"{activity['name']}-签到结果.csv"
    return Response(
        "\ufeff" + output.getvalue(),
        mimetype="text/csv; charset=utf-8",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )


@app.errorhandler(413)
def too_large(_error):
    return "上传文件不能超过 8 MB。", 413


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", 5000)), debug=os.environ.get("FLASK_DEBUG") == "1")
